import openpyxl
from openpyxl.drawing.image import Image
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
import xlsxwriter

if "__main__" == __name__:

    # Abrimos dataframes
    comprar = pd.read_csv("compra_semanal.csv", sep =',')
    orders = pd.read_csv("order_details_ordenado.csv", sep =',')
    pizzas = pd.read_csv('pizzas.csv')

    # Los ordenamos
    pizzas = pizzas.sort_values('price', ascending=False)

    with pd.ExcelWriter("Informe_Pizzerias_Maven.xlsx", mode="w", engine="xlsxwriter") as writer:
        pizzas.to_excel(writer, sheet_name='Reporte ejecutivo', index=False)
        comprar.to_excel(writer, sheet_name='Reporte ingredientes', index=False)
        orders.to_excel(writer, sheet_name='Reporte pedidos', index=False)

        wb  = writer.book

        # Gráficos primera hoja
        hoja1 = 'Reporte ejecutivo'

        chart1 = wb.add_chart({'type': 'bar'})
        chart1.add_series({'categories': f'{hoja1}!$B$2:$B$6', 'values': f'{hoja1}!$D$2:$D$6'})

        chart1.set_title ({'name': 'Pizzas más caras'})
        chart1.set_x_axis({'name': 'Euros'})
        chart1.set_y_axis({'name': 'Pizzas'})
        chart1.set_style(11)

        writer.sheets['Reporte ejecutivo'].insert_chart('F2', chart1, {'x_offset': 25, 'y_offset': 10})

        chart2 = wb.add_chart({'type': 'bar'})
        chart2.add_series({'categories': f'{hoja1}!$B$93:$B$97', 'values': f'{hoja1}!$D$93:$D$97'})

        chart2.set_title ({'name': 'Pizzas más baratas'})
        chart2.set_x_axis({'name': 'Euros'})
        chart2.set_y_axis({'name': 'Pizzas'})
        chart2.set_style(11)

        writer.sheets['Reporte ejecutivo'].insert_chart('F22', chart2, {'x_offset': 25, 'y_offset': 10})

        # Gráficos segunda hoja
        hoja2 = 'Reporte ingredientes'

        chart3 = wb.add_chart({'type': 'bar'})
        chart3.add_series({'categories': f'{hoja2}!$A$2:$A$66', 'values': f'{hoja2}!$B$2:$B$66'})

        chart3.set_title ({'name': 'Ingredientes'})
        chart3.set_x_axis({'name': 'Porciones'})
        chart3.set_y_axis({'name': 'Ingredientes'})
        chart3.set_style(11)

        writer.sheets['Reporte ingredientes'].insert_chart('F2', chart3, {'x_offset': 50, 'y_offset': 20})
